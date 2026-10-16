import openpyxl
import os

file = "atm_data.xlsx"

# Create Excel file if it does not exist
if not os.path.exists(file):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Name", "Phone Number", "Balance"])
    wb.save(file)


def add_new():
    name = input("Enter your name: ")
    phone = input("Enter phone number: ")
    balance = float(input("Enter initial deposit: "))

    wb = openpyxl.load_workbook(file)
    sheet = wb.active
    sheet.append([name, phone, balance])
    wb.save(file)

    print("Account created successfully!")


def deposit():
    phone = input("Enter phone number: ")
    amount = float(input("Enter amount to deposit: "))

    wb = openpyxl.load_workbook(file)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[1].value) == phone:
            if amount > 0:
                row[2].value += amount
                wb.save(file)
                print("Money deposited successfully!")
            else:
                print("Enter a valid amount.")
            return

    print("Account not found.")


def withdraw():
    phone = input("Enter phone number: ")

    wb = openpyxl.load_workbook(file)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[1].value) == phone:
            amount = float(input("Enter amount to withdraw: "))

            if amount <= 0:
                print("Enter a valid amount.")
            elif amount <= row[2].value:
                row[2].value -= amount
                wb.save(file)
                print("Money withdrawn successfully!")
            else:
                print("Insufficient balance.")
            return

    print("Account not found.")


def view_balance():
    phone = input("Enter phone number: ")

    wb = openpyxl.load_workbook(file)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[1].value) == phone:
            print("Account holder:", row[0].value)
            print("Current balance: ₹", row[2].value)
            return

    print("Account not found.")


def menu():
    while True:
        print("\n===== WELCOME TO ATM SYSTEM =====")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. View Balance")
        print("5. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_new()
        elif choice == "2":
            deposit()
        elif choice == "3":
            withdraw()
        elif choice == "4":
            view_balance()
        elif choice == "5":
            print("Thank you for using our ATM!")
            break
        else:
            print("Invalid choice. Try again.")


menu()
